"""
Builds public/templates/teams-template.xlsx, the file admins download from the
setup page to upload teams with their tracks.

    python3 -m pip install openpyxl
    python3 scripts/make-teams-template.py

The first sheet MUST stay the teams sheet: the app reads only the first sheet.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

MAROON = "80191A"
GOLD = "F9D69F"
CREAM = "FDEFD9"
FONT = "Arial"

OUT = Path(__file__).resolve().parent.parent / "public" / "templates" / "teams-template.xlsx"

EXAMPLES = [
    ("فريق الأمل", "الصحة"),
    ("فريق النور", "الصحة"),
    ("فريق الشفاء", "الصحة"),
    ("فريق المعرفة", "التعليم"),
    ("فريق الإبداع", "التعليم"),
    ("فريق المستقبل", "التعليم"),
    ("فريق الأرض الخضراء", "الاستدامة"),
    ("فريق المياه", "الاستدامة"),
    ("فريق الطاقة", "الاستدامة"),
    ("فريق الرقمنة", "التقنية"),
    ("فريق البيانات", "التقنية"),
    ("فريق الذكاء", "التقنية"),
]

INSTRUCTIONS = [
    "اكتب كل فريق في سطر واحد: اسم الفريق في العمود A، والمسار في العمود B.",
    "لا تغيّر سطر العناوين (السطر الأول): «اسم الفريق» و«المسار».",
    "اسم الفريق مطلوب، ولا يتكرر في الملف.",
    "المسار مطلوب لكل فريق. اكتب اسم المسار بنفس الطريقة لكل فرقه، مثلاً «الصحة» لكل فرق مسار الصحة.",
    "الأسطر الموجودة أمثلة فقط: احذفها واكتب فرقك.",
    "ترتيب الفرق في الملف هو ترتيب التحكيم.",
    "احفظ الملف بصيغة Excel ‏(.xlsx). ملفات CSV مقبولة أيضاً.",
    "صيغة xls القديمة وملفات Numbers غير مدعومة. من تطبيق Numbers على الآيباد: مشاركة ← تصدير ← Excel.",
    "يُقرأ أول ورقة في الملف فقط (ورقة «الفرق»)، فلا تغيّر ترتيب الأوراق.",
    "الحد الأقصى: 500 فريق، و80 حرفاً لاسم الفريق، و60 حرفاً لاسم المسار.",
    "رفع الملف مرة أخرى لا يكرر الفرق: الفريق الموجود بنفس الاسم يُحدَّث مساره وترتيبه فقط.",
]


def build() -> None:
    wb = Workbook()

    # ---------------------------------------------------------------- teams --
    ws = wb.active
    ws.title = "الفرق"
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    thin = Side(style="thin", color="D9C6B0")

    for col, title in (("A", "اسم الفريق"), ("B", "المسار")):
        c = ws[f"{col}1"]
        c.value = title
        c.font = Font(name=FONT, bold=True, size=13, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MAROON)
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[1].height = 26

    for i, (name, track) in enumerate(EXAMPLES, start=2):
        for col, value in (("A", name), ("B", track)):
            c = ws[f"{col}{i}"]
            c.value = value
            c.font = Font(name=FONT, size=12)
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.border = Border(bottom=thin)
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=CREAM)

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 28

    # Gentle guardrails inside Excel itself (the app validates again on upload)
    name_rule = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="80", allow_blank=True)
    name_rule.error = "اسم الفريق يجب ألا يزيد عن 80 حرفاً"
    name_rule.errorTitle = "اسم طويل"
    name_rule.prompt = "اسم الفريق: مطلوب ولا يتكرر"
    name_rule.promptTitle = "اسم الفريق"
    track_rule = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="60", allow_blank=True)
    track_rule.error = "اسم المسار يجب ألا يزيد عن 60 حرفاً"
    track_rule.errorTitle = "اسم طويل"
    track_rule.prompt = "المسار: اكتبه بنفس الطريقة لكل فرقه"
    track_rule.promptTitle = "المسار"
    for rule, rng in ((name_rule, "A2:A501"), (track_rule, "B2:B501")):
        rule.showErrorMessage = True
        rule.showInputMessage = True
        ws.add_data_validation(rule)
        rule.add(rng)

    # --------------------------------------------------------- instructions --
    info = wb.create_sheet("تعليمات")
    info.sheet_view.rightToLeft = True
    info["A1"] = "طريقة تعبئة ملف الفرق"
    info["A1"].font = Font(name=FONT, bold=True, size=15, color=MAROON)
    info.row_dimensions[1].height = 28
    for i, line in enumerate(INSTRUCTIONS, start=3):
        info[f"A{i}"] = f"{i - 2}. {line}"
        info[f"A{i}"].font = Font(name=FONT, size=12)
        info[f"A{i}"].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        info.row_dimensions[i].height = 34
    info.column_dimensions["A"].width = 110

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
